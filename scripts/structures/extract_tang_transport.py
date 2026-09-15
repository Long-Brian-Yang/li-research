"""Extract selected public workbook arrays; no resampling or curve alteration.

S24 column headers are reversed relative to the published figure axes:
first column is time (ps), second is MSD (angstrom squared), verified on SI p19.
S3 vertical axis is ln[sigma(S/cm)*T(K)], verified on SI p4.
"""
from pathlib import Path
import csv,json,hashlib,math
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'results/amorphous_review_20260915'
src=BASE/'source/literature/Tang2026_source.xlsx'
out=BASE/'completed_transport';out.mkdir(exist_ok=True)
w=openpyxl.load_workbook(src,read_only=True,data_only=True)
def save(name,header,rows):
    with (out/name).open('w') as f:
        writer=csv.writer(f,lineterminator='\n');writer.writerow(header);writer.writerows(rows)
for j,T in enumerate([320,330,340,350]):
    rows=[r[2*j:2*j+2] for r in w['S24'].iter_rows(min_row=4,max_col=8,values_only=True)
          if all(isinstance(v,(int,float)) for v in r[2*j:2*j+2])]
    save(f'Tang_S24_{T}K.csv',['time_ps','MSD_A2'],rows)
rows=[]
for r in w['Figure 3'].iter_rows(min_row=3,max_col=8,values_only=True):
    x,y,e=r[5:8]
    if not isinstance(x,(int,float)):continue
    T=1000/x;rows.append([x,T,y,e,math.exp(y)/T*1000])
save('Tang_Fig3g.csv',['inverse_T_1000_K','T_from_rounded_x_K','ln_sigmaT_S_cm_K','reported_y_error','sigma_mS_cm'],rows)
rows=[]
for r in w['S3'].iter_rows(min_row=3,max_col=4,values_only=True):
    x,y=r[0],r[3]
    if not isinstance(x,(int,float)):continue
    T=1000/x;rows.append([x,T,y,math.exp(y)/T*1000])
save('Tang_S3_experiment.csv',['inverse_T_1000_K','T_from_rounded_x_K','ln_sigmaT_S_cm_K','sigma_mS_cm'],rows)
(out/'Tang_extraction.json').write_text(json.dumps({'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),
    'doi':'10.1038/s41467-026-69737-x','MSD_axes':'S24 time first, MSD second; workbook headers reversed, SI p19 controls',
    'experimental_axis':'SI p4 S3: ln(sigma T), sigma in S/cm',
    'Figure3g':'MLFF-MD: 320–350 K 300ps, ~300K separate 3ns; x rounded by source; error definition not independently established',
    'conversion':'sigma(mS/cm)=1000*exp(y)/T; T=1000/x'},indent=2)+'\n')
print('Extracted 6 source tables')
