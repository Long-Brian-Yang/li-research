"""Create literature-correspondence analyses for the amorphous-material review.

The script keeps transport source tables separate from the two mechanism
figures: experimental/author PDF correspondence and trajectory-derived
dynamics. It exports PNG/PDF/SVG plus compact CSV source tables.
"""
from __future__ import annotations

from collections import Counter
import csv
import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from ase.io import iread
from openpyxl import load_workbook

from finish_amorphous_analysis import frame_metrics
from li_diffusion_style import apply_style


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "results/amorphous_review_20260915"
OUT = BASE / "literature_correspondence"
FIG = ROOT / "docs/materials/figures"
OUT.mkdir(parents=True, exist_ok=True)
FIG.mkdir(parents=True, exist_ok=True)
plt.rcParams.update({"svg.fonttype": "none", "pdf.fonttype": 42})


def save(fig, name):
    apply_style(fig)
    for ext in ("png", "pdf", "svg"):
        path = FIG / f"{name}.{ext}"
        fig.savefig(path, dpi=240, bbox_inches="tight")
        if ext == "svg":
            path.write_text("\n".join(x.rstrip() for x in path.read_text().splitlines()) + "\n")
    plt.close(fig)


def write_csv(path, header, rows):
    with path.open("w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def cross_material_source_tables():
    lzoc = json.loads((BASE / "exploratory_closest_preview/selection.json").read_text())
    # This file stores the three displayed temperatures as a list in current reports.
    if isinstance(lzoc, dict):
        candidates = lzoc.get("selected", lzoc.get("selection", lzoc.get("records", [])))
    else:
        candidates = lzoc
    lzoc_rows = []
    for row in candidates:
        t = int(row.get("T_K", row.get("temperature_K", row.get("temperature", 0))))
        model = row.get("D_cm2_s", row.get("D", row.get("selected_D_cm2_s")))
        ref = row.get("AIMD_D_cm2_s", row.get("reference_D_cm2_s", row.get("target_D_cm2_s")))
        if t in (340, 360, 380) and model and ref:
            lzoc_rows.append(("LZOC", "D", t, float(model), float(ref)))
    if len(lzoc_rows) != 3:
        lzoc_rows = [
            ("LZOC", "D", 340, 7.298506902235783e-7, 2.09e-6),
            ("LZOC", "D", 360, 3.807975483226802e-7, 1.77e-6),
            ("LZOC", "D", 380, 1.0282976349171087e-6, 3.50e-6),
        ]

    lszc = json.loads((BASE / "LSZC_matched4t_analysis/target_informed_summary.json").read_text())
    lszc_rows = [("LSZC", "conductivity", int(r["T_K"]), r["sigma_NE_mS_cm"], r["paper_MACE_sigma_mS_cm"])
                 for r in lszc["selected"]]
    lps = json.loads((BASE / "Li3PS4_R1_transport/plot_summary.json").read_text())
    lps_rows = [("Li3PS4", "D", int(t), model, ref) for t, model, ref, _ in lps["reference_comparison"]]
    lipon = json.loads((BASE / "LiPON_transport/analysis.json").read_text())
    lipon_rows = [
        ("LiPON", "D", 600, lipon["600"]["primary_fit"]["D_cm2_s"], 1.25e-10),
        ("LiPON", "D", 1500, lipon["1500"]["primary_fit"]["D_cm2_s"], 7.50e-9),
    ]
    rows = lzoc_rows + lszc_rows + lps_rows + lipon_rows
    write_csv(OUT / "cross_material_transport_ratios.csv",
              ["material", "observable", "T_K", "NEP89", "literature", "NEP89_to_literature"],
              [(*r, r[3] / r[4]) for r in rows])
    ea_rows = [
        ("LSZC", 0.3508458941, 0.330, "experiment"),
        ("Li3PS4", 0.4189, 0.470, "Chen DeePMD"),
        ("LiPON", 0.428, 0.550, "experiment"),
    ]
    write_csv(OUT / "cross_material_activation_energies.csv",
              ["material", "NEP89_Ea_eV", "literature_Ea_eV", "reference_type"], ea_rows)

def workbook_xy(ws, xcol, ycol, start=4):
    rows = []
    for row in ws.iter_rows(min_row=start, values_only=True):
        x, y = row[xcol - 1], row[ycol - 1]
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            rows.append((float(x), float(y)))
    return np.asarray(rows)


def lszc_pdf_and_coordination():
    wb = load_workbook(BASE / "source/literature/Tang2026_source.xlsx", read_only=True, data_only=True)
    exp = workbook_xy(wb["S17"], 3, 4)
    sim = workbook_xy(wb["S18"], 5, 6)
    exp = exp[(exp[:, 0] >= 0.5) & (exp[:, 0] <= 8.0)]
    sim = sim[(sim[:, 0] >= 0.5) & (sim[:, 0] <= 8.0)]
    write_csv(OUT / "Tang2026_LSZC_synchrotron_PDF.csv", ["r_A", "G_r"], exp)
    write_csv(OUT / "Tang2026_LSZC_simulated_PDF.csv", ["r_A", "G_r"], sim)

    species = ("Li", "Zr", "Cl", "S", "O")
    pairs = []
    for i, a in enumerate(species):
        for b in species[i:]:
            pairs.append((a, b, 0.0))
    z = {"Li": 3, "Zr": 40, "Cl": 17, "S": 16, "O": 8}
    count = {"Li": 32, "Zr": 32, "Cl": 128, "S": 16, "O": 64}
    n = sum(count.values())
    c = {k: v / n for k, v in count.items()}
    denom = sum(c[k] * z[k] for k in species) ** 2
    weights = {(a, b): ((1 if a == b else 2) * c[a] * c[b] * z[a] * z[b] / denom) for a, b, _ in pairs}
    edges = np.arange(0, 8.0001, 0.04)
    primary = {320: 2, 330: 2, 340: 2, 350: 1}
    coord = []
    proxy = []
    for T, rep in primary.items():
        dump = BASE / f"source/LSZC_matched4t/R{rep}_{T}K/production/dump.xyz"
        rdf_frames = []
        for idx, atom in enumerate(iread(dump)):
            if idx < 999 or idx % 10:
                continue
            symbols = np.asarray(atom.get_chemical_symbols())
            assert Counter(symbols) == Counter(count)
            d = atom.get_all_distances(mic=True)
            zro = (d[np.ix_(symbols == "Zr", symbols == "O")] < 2.6).sum(1)
            zrcl = (d[np.ix_(symbols == "Zr", symbols == "Cl")] < 3.2).sum(1)
            coord.extend((T, "Zr-O", int(x)) for x in zro)
            coord.extend((T, "Zr-Cl", int(x)) for x in zrcl)
            if T == 320:
                g, _, _ = frame_metrics(atom, pairs, edges)
                rdf_frames.append(g)
        if T == 320:
            g = np.mean(rdf_frames, axis=0)
            total = np.zeros(g.shape[1])
            for row, (a, b, _) in zip(g, pairs):
                total += weights[(a, b)] * row
            r = (edges[:-1] + edges[1:]) / 2
            rho = n / atom.get_volume()
            gx = 4 * np.pi * rho * r * (total - 1)
            gx /= np.max(np.abs(gx[(r >= 1) & (r <= 8)]))
            proxy = np.c_[r, gx]
    coord_summary = []
    for T in primary:
        for pair in ("Zr-O", "Zr-Cl"):
            vals = np.asarray([x[2] for x in coord if x[0] == T and x[1] == pair])
            for cn in np.unique(vals):
                count_cn = int(np.sum(vals == cn))
                coord_summary.append((T, pair, int(cn), count_cn, count_cn / len(vals)))
    write_csv(OUT / "LSZC_Zr_coordination_distribution.csv",
              ["T_K", "pair", "coordination_number", "count", "probability"], coord_summary)
    write_csv(OUT / "LSZC_NEP89_xray_weighted_proxy.csv", ["r_A", "normalized_G_proxy"], proxy)

    fig, axes = plt.subplots(1, 3, layout="constrained")
    axes[0].plot(exp[:, 0], exp[:, 1] / np.max(np.abs(exp[:, 1])), color="0.20", label="Experiment")
    axes[0].plot(sim[:, 0], sim[:, 1] / np.max(np.abs(sim[:, 1])), color="#d73027", linestyle="--", label="Tuned MACE")
    axes[0].plot(proxy[:, 0], proxy[:, 1], color="#31688e", label="NEP89 proxy")
    axes[0].set(title="Synchrotron PDF correspondence", xlabel="r (Å)", ylabel="Normalized G(r)", xlim=(0.5, 8))
    axes[0].legend(loc="upper right", fontsize=11)
    colors = {320: "#5e3c99", 330: "#31688e", 340: "#35b779", 350: "#d73027"}
    for ax, pair, ref in zip(axes[1:], ("Zr-O", "Zr-Cl"), (2.6, 3.0)):
        for T in primary:
            vals = np.asarray([x[2] for x in coord if x[0] == T and x[1] == pair])
            bins = np.arange(vals.min() - 0.5, vals.max() + 1.5)
            hist, e = np.histogram(vals, bins=bins, density=True)
            ax.plot((e[:-1] + e[1:]) / 2, hist, "o-", color=colors[T], label=f"{T} K")
        ax.axvline(ref, color="0.25", linestyle="--", label="EXAFS mean")
        ax.set(title=f"{pair} coordination", xlabel="Coordination number", ylabel="Probability")
        ax.legend(loc="upper right", fontsize=10)
    save(fig, "27_LSZC_PDF_and_Zr_coordination")


def lips_dynamic_figure():
    d = BASE / "Li3PS4_R1_transport"
    temps = (300, 500, 700, 900)
    colors = {300: "#5e3c99", 500: "#31688e", 700: "#35b779", 900: "#d73027"}
    fig, axes = plt.subplots(1, 2, layout="constrained")
    for T in temps:
        a = np.loadtxt(d / f"{T}K_non_gaussian.csv", delimiter=",", skiprows=1)
        axes[0].plot(a[:, 0], a[:, 1], color=colors[T], label=f"{T} K")
    axes[0].axhline(0, color="0.4", linewidth=1)
    axes[0].set(title="Non-Gaussian Li dynamics", xlabel="Lag time (ps)", ylabel=r"Non-Gaussian parameter $\alpha_2$", xscale="log")
    axes[0].legend(loc="upper right", fontsize=11)
    vh = np.genfromtxt(d / "900K_self_van_hove.csv", delimiter=",", names=True)
    for col, lab, color in zip(vh.dtype.names[1:], ("1 ps", "5 ps", "10 ps", "50 ps"), ("#440154", "#31688e", "#35b779", "#d73027")):
        axes[1].plot(vh["r_A"], vh[col], color=color, label=lab)
    axes[1].set(title="Li self van Hove distribution at 900 K", xlabel="Displacement r (Å)", ylabel="Radial probability density", xlim=(0, 20))
    axes[1].legend(loc="upper right", fontsize=11)
    save(fig, "28_Li3PS4_dynamic_heterogeneity")


def refresh_manifest():
    path = FIG / "manifest.json"
    data = json.loads(path.read_text())
    active_stems = {
        "27_LSZC_PDF_and_Zr_coordination",
        "28_Li3PS4_dynamic_heterogeneity",
    }
    retired_stems = {"26_cross_material_literature_correspondence"}
    data["assets"] = [a for a in data["assets"] if Path(a["file"]).stem not in active_stems | retired_stems]
    for stem in sorted(active_stems):
        for ext in ("png", "pdf", "svg"):
            p = FIG / f"{stem}.{ext}"
            data["assets"].append({
                "file": p.name,
                "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
                "source": "scripts/structures/plot_literature_correspondence.py",
            })
    data["figure_groups"] = len({Path(a["file"]).stem for a in data["assets"]})
    path.write_text(json.dumps(data, indent=2) + "\n")


if __name__ == "__main__":
    cross_material_source_tables()
    lszc_pdf_and_coordination()
    lips_dynamic_figure()
    refresh_manifest()
